
"""
ApexPlanet Data Analytics Internship — Task 5
Automated Data Pipeline Script
Runs: Load → Clean → Analyse → Export to Excel → Save Charts
"""

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import os
import warnings
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

warnings.filterwarnings("ignore")

print("=" * 55)
print("   APEXPLANET — AUTOMATED DATA PIPELINE")
print(f"   Run Time: {datetime.now().strftime(\%Y-%m-%d %H:%M:%S\)}")
print("=" * 55)

# ── STEP 1: LOAD RAW DATA ──
print("\n[1/5] Loading raw data...")
df = pd.read_csv("../data/data.csv", encoding="ISO-8859-1")
print(f"      Loaded: {len(df):,} rows")

# ── STEP 2: CLEAN DATA ──
print("[2/5] Cleaning data...")
df.drop_duplicates(inplace=True)
df.dropna(subset=["CustomerID"], inplace=True)
df["Description"].fillna("Unknown", inplace=True)
df["InvoiceDate"] = pd.to_datetime(df["InvoiceDate"])
df["CustomerID"]  = df["CustomerID"].astype(int).astype(str)
df["Year"]  = df["InvoiceDate"].dt.year
df["Month"] = df["InvoiceDate"].dt.month
df["Day"]   = df["InvoiceDate"].dt.day
df = df[(df["Quantity"] > 0) & (df["UnitPrice"] > 0)]
for col in ["Quantity", "UnitPrice"]:
    Q1, Q3 = df[col].quantile([0.25, 0.75])
    IQR = Q3 - Q1
    df = df[(df[col] >= Q1 - 1.5*IQR) & (df[col] <= Q3 + 1.5*IQR)]
df["TotalPrice"] = df["Quantity"] * df["UnitPrice"]
df.to_csv("../data/data_cleaned_auto.csv", index=False)
print(f"      Clean rows: {len(df):,} | Saved → data/data_cleaned_auto.csv")

# ── STEP 3: COMPUTE KPIs ──
print("[3/5] Computing KPIs...")
kpis = {
    "Total Revenue (£)":    round(df["TotalPrice"].sum(), 2),
    "Total Orders":         df["InvoiceNo"].nunique(),
    "Unique Customers":     df["CustomerID"].nunique(),
    "Unique Products":      df["Description"].nunique(),
    "Avg Order Value (£)":  round(df["TotalPrice"].sum() / df["InvoiceNo"].nunique(), 2),
    "Top Country":          df.groupby("Country")["TotalPrice"].sum().idxmax(),
}
for k, v in kpis.items():
    print(f"      {k}: {v}")

# ── STEP 4: EXPORT TO EXCEL ──
print("[4/5] Exporting to Excel...")
monthly  = df.groupby(["Year","Month"]).agg(
    Revenue=("TotalPrice","sum"), Orders=("InvoiceNo","nunique"),
    Customers=("CustomerID","nunique")
).reset_index().round(2)

countries = df.groupby("Country").agg(
    Revenue=("TotalPrice","sum"), Customers=("CustomerID","nunique")
).sort_values("Revenue",ascending=False).head(10).reset_index().round(2)

products = df.groupby("Description").agg(
    Revenue=("TotalPrice","sum"), Quantity=("Quantity","sum")
).sort_values("Revenue",ascending=False).head(10).reset_index().round(2)

wb = Workbook()
header_font  = Font(bold=True, color="FFFFFF", size=11)
header_fill  = PatternFill("solid", fgColor="1565C0")
center_align = Alignment(horizontal="center", vertical="center")
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)

def write_sheet(wb, title, df_data, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(list(df_data.columns))
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment, cell.border = center_align, thin_border
    for row in dataframe_to_rows(df_data, index=False, header=False):
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

# KPI sheet
ws_kpi = wb.active
ws_kpi.title = "KPI Summary"
ws_kpi.append(["Metric", "Value"])
for cell in ws_kpi[1]:
    cell.font, cell.fill = header_font, header_fill
    cell.alignment = center_align
for k, v in kpis.items():
    ws_kpi.append([k, v])
for col in ws_kpi.columns:
    ws_kpi.column_dimensions[col[0].column_letter].width = 30

write_sheet(wb, "Monthly Revenue", monthly)
write_sheet(wb, "Top Countries",   countries)
write_sheet(wb, "Top Products",    products)

wb.save("../reports/analytics_report.xlsx")
print("      Saved → reports/analytics_report.xlsx")

# ── STEP 5: SAVE SUMMARY CHARTS ──
print("[5/5] Generating summary charts...")
monthly["Period"] = pd.to_datetime(
    monthly["Year"].astype(str) + "-" + monthly["Month"].astype(str).str.zfill(2)
)

fig, axes = plt.subplots(1, 2, figsize=(14, 5))
axes[0].plot(monthly["Period"], monthly["Revenue"],
             color="#1565C0", marker="o", linewidth=2)
axes[0].set_title("Monthly Revenue Trend")
axes[0].set_ylabel("Revenue (£)")
axes[0].tick_params(axis="x", rotation=45)

axes[1].barh(countries["Country"], countries["Revenue"],
             color="#42A5F5")
axes[1].set_title("Top 10 Countries by Revenue")
axes[1].set_xlabel("Revenue (£)")

plt.suptitle("Automated Pipeline — Summary Charts", fontsize=14, fontweight="bold")
plt.tight_layout()
plt.savefig("../reports/pipeline_summary_charts.png", dpi=150, bbox_inches="tight")
plt.close()
print("      Saved → reports/pipeline_summary_charts.png")

print("\n" + "=" * 55)
print("   ✅ PIPELINE COMPLETE!")
print("=" * 55)
